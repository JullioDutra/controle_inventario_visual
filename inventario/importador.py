"""Leitura e aplicação de exportações do TopDesk (xlsx/csv).

Regras:
- A chave é o Asset ID. Existentes são atualizados, novos criados, ausentes sinalizados (nunca apagados).
- Só são atualizadas as colunas que EXISTEM no arquivo; uma aba sem "Fornecedor" não apaga o fornecedor.
- Campos locais (setor, pavimento, posição na planta) nunca são sobrescritos, exceto o setor quando ainda vazio.
- Colunas desconhecidas vão para `extras` (modelo, marca, ...), então a planilha completa não perde dados.
"""
import csv
import io
import re
import unicodedata

from django.db import transaction

from .models import Equipamento, Historico, Setor

TIPOS = dict(Equipamento.Tipo.choices)
STATUS = dict(Equipamento.Status.choices)

# Setores inferidos do campo Obs (só quando o equipamento ainda não tem setor).
REGRAS_SETOR = [
    (r"kiosk\s+produ", "Produção"), (r"kiosk\s+lcq|laborat|lcq", "Laboratório (LCQ)"),
    (r"kiosk\s+alimentos|smac", "Alimentos / SMAC"), (r"manuten", "Manutenção"),
    (r"engenharia", "Engenharia"), (r"descanso", "Sala de descanso"),
    (r"portaria|camera", "Portaria / Segurança"), (r"seguran[cç]a do trabalho", "SST"),
]
COLUNAS = {
    "asset_id": ["asset id", "assetid", "id", "asset", "ativo", "patrimonio", "id do ativo"],
    "obs": ["obs", "observacao", "observacoes", "notes", "remarks"],
    "serial": ["serial number", "serial", "numero de serie", "n/s", "serie", "serial no"],
    "status": ["status", "situacao", "estado"],
    "fornecedor": ["supplier", "fornecedor"],
    "responsavel": ["assignments - persons", "assignments", "assigned to", "pessoa", "pessoas", "usuario",
                    "responsavel", "colaborador"],
    "tipo": ["tipo", "type", "categoria", "category", "asset type"],
    "setor": ["setor", "departamento", "department"],
}
INDICE = {sinonimo: campo for campo, lista in COLUNAS.items() for sinonimo in lista}


def norm(s):
    s = unicodedata.normalize("NFD", str(s if s is not None else ""))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").strip().lower()


def texto(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def tipo_por_texto(nome, asset_id=""):
    n, p = norm(nome), norm(asset_id)[:3]
    if re.search(r"laptop|notebook", n):
        return "laptop"
    if re.search(r"tiny|\bpc\b|desktop|computador", n):
        return "tinypc"
    if re.search(r"impress|printer|multifuncional", n):
        return "impressora"
    return {"plp": "laptop", "ppc": "tinypc", "pim": "impressora", "pip": "impressora"}.get(p, "outro")


def status_norm(s):
    n = norm(s)
    if re.search(r"operat|operac", n):
        return "operacional"
    if re.search(r"backup|reserva", n):
        return "backup"
    if re.search(r"install|instalar", n):
        return "instalar"
    return "outro"


def setor_por_obs(obs):
    for rx, nome in REGRAS_SETOR:
        if re.search(rx, obs or "", re.I):
            return nome
    return ""


def _linhas_da_tabela(cabecalho, corpo, nome_aba):
    cols = [(i, norm(h)) for i, h in enumerate(cabecalho) if texto(h)]
    campos = {i: INDICE.get(n) for i, n in cols}
    if "asset_id" not in campos.values():
        return []
    presentes = sorted({c for c in campos.values() if c})
    saida = []
    for r in corpo:
        d, extras = {}, {}
        for i, h in enumerate(cabecalho):
            if i >= len(r) or not texto(h):
                continue
            v = texto(r[i])
            campo = campos.get(i)
            if campo:
                d[campo] = v
            elif v:
                extras[texto(h)] = v
        aid = d.get("asset_id", "").upper()
        if not aid:
            continue
        tipo = tipo_por_texto(d["tipo"], aid) if d.get("tipo") else tipo_por_texto(nome_aba, aid)
        saida.append({
            "asset_id": aid, "tipo": tipo, "status": status_norm(d.get("status", "")) if "status" in presentes else None,
            "serial": d.get("serial", ""), "fornecedor": d.get("fornecedor", ""), "responsavel": d.get("responsavel", ""),
            "obs": d.get("obs", ""), "setor": d.get("setor", ""), "extras": extras, "cols": presentes, "aba": nome_aba,
        })
    return saida


def ler_arquivo(nome, conteudo):
    """Devolve (linhas, abas). `conteudo` = bytes."""
    linhas, abas = [], []
    if nome.lower().endswith(".csv"):
        txt = conteudo.decode("utf-8-sig", errors="replace")
        try:
            dialeto = csv.Sniffer().sniff(txt[:4096], delimiters=",;\t")
        except csv.Error:
            dialeto = csv.excel
        rows = list(csv.reader(io.StringIO(txt), dialeto))
        if rows:
            l = _linhas_da_tabela(rows[0], rows[1:], nome.rsplit(".", 1)[0])
            if l:
                abas.append({"nome": nome, "n": len(l)})
                linhas += l
        return linhas, abas
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        l = _linhas_da_tabela(list(rows[0]), [list(r) for r in rows[1:]], ws.title)
        if l:
            abas.append({"nome": ws.title, "n": len(l)})
            linhas += l
    return linhas, abas


CAMPOS = [("tipo", "Tipo"), ("status", "Status"), ("responsavel", "Responsável"), ("serial", "Nº série"),
          ("fornecedor", "Fornecedor"), ("obs", "Obs")]


def _fmt(campo, v):
    if campo == "status":
        return STATUS.get(v, v) or "—"
    if campo == "tipo":
        return TIPOS.get(v, v) or "—"
    return v or "—"


def processar(linhas, arquivo, aplicar=False):
    """Calcula (e, se `aplicar`, grava) a importação. Devolve (resumo, mudancas)."""
    res = {"novos": 0, "alterados": 0, "iguais": 0, "ausentes": 0, "duplicados": 0, "voltaram": 0}
    mud = []
    existentes = {e.asset_id: e for e in Equipamento.objects.all()}
    setores = {s.nome.lower(): s for s in Setor.objects.all()}
    vistos, historicos, criar, atualizar = set(), [], [], []

    def setor_obj(nome):
        if not nome:
            return None
        s = setores.get(nome.lower())
        if s is None and aplicar:
            s = Setor.objects.create(nome=nome)
            setores[nome.lower()] = s
        return s

    def log(aid, ev, det):
        mud.append({"id": aid, "evento": ev, "detalhe": det})
        historicos.append(Historico(asset_id=aid, evento=ev, detalhe=det, origem=arquivo))

    for n in linhas:
        aid = n["asset_id"]
        if aid in vistos:
            res["duplicados"] += 1
            continue
        vistos.add(aid)
        e = existentes.get(aid)
        nome_setor = n["setor"] or setor_por_obs(n["obs"])
        if e is None:
            res["novos"] += 1
            log(aid, "novo", f"Novo {TIPOS[n['tipo']]} ({STATUS.get(n['status'] or 'operacional')})")
            if aplicar:
                criar.append(Equipamento(
                    asset_id=aid, tipo=n["tipo"], status=n["status"] or "operacional", serial=n["serial"],
                    fornecedor=n["fornecedor"], responsavel=n["responsavel"], obs=n["obs"], extras=n["extras"],
                    setor=setor_obj(nome_setor)))
            continue
        difs, mudou = [], False
        for campo, rotulo in CAMPOS:
            if campo not in n["cols"] and campo != "tipo":
                continue
            novo = n[campo]
            if campo == "tipo" and (novo == "outro" or novo == e.tipo):
                continue
            if campo == "status" and novo is None:
                continue
            antigo = getattr(e, campo)
            if (antigo or "") != (novo or ""):
                difs.append(f"{rotulo}: {_fmt(campo, antigo)} → {_fmt(campo, novo)}")
                setattr(e, campo, novo)
                mudou = True
        for k, v in n["extras"].items():
            if e.extras.get(k) != v:
                difs.append(f"{k}: {e.extras.get(k, '—')} → {v}")
                e.extras = {**e.extras, k: v}
                mudou = True
        if e.setor_id is None and nome_setor:
            s = setor_obj(nome_setor)
            if s is not None or not aplicar:
                if aplicar:
                    e.setor = s
                difs.append(f"Setor: — → {nome_setor}")
                mudou = True
        if e.ausente:
            e.ausente = False
            difs.append("Voltou a aparecer no TopDesk")
            res["voltaram"] += 1
            mudou = True
        if difs:
            res["alterados"] += 1
            log(aid, "alterado", " · ".join(difs))
        else:
            res["iguais"] += 1
        if mudou:
            atualizar.append(e)
    tipos_no_arquivo = {n["tipo"] for n in linhas}
    for aid, e in existentes.items():
        if aid not in vistos and e.tipo in tipos_no_arquivo and not e.ausente:
            res["ausentes"] += 1
            log(aid, "ausente", "Não veio na última exportação do TopDesk")
            e.ausente = True
            atualizar.append(e)
    if aplicar:
        with transaction.atomic():
            Equipamento.objects.bulk_create(criar)
            if atualizar:
                Equipamento.objects.bulk_update(
                    atualizar, ["tipo", "status", "responsavel", "serial", "fornecedor", "obs", "extras", "setor", "ausente"])
            Historico.objects.bulk_create(historicos)
    return res, mud
